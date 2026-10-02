#!/usr/bin/env python3
import argparse
import sqlite3
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Importa a planilha INTEGRACAO_OK_QLIK_ARIUS e marca devolucoes como concluidas."
    )
    parser.add_argument(
        "--db",
        default="outputs/devolucao_auditoria.sqlite",
        help="Caminho do banco SQLite.",
    )
    parser.add_argument(
        "--input",
        default="/Users/alexandrematheuscrose/Downloads/INTEGRACAO_OK_QLIK_ARIUS.xlsx",
        help="Caminho do arquivo XLSX.",
    )
    parser.add_argument(
        "--sheet",
        default="COMPARA",
        help="Nome da aba com os registros OK.",
    )
    return parser.parse_args()


def ensure_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS devolucao_export_lote (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            file_name TEXT NOT NULL,
            file_path TEXT NOT NULL UNIQUE,
            exported_at TEXT NOT NULL,
            total_items INTEGER NOT NULL DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS devolucao_export_item (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            lote_id INTEGER NOT NULL,
            nro_loja INTEGER NOT NULL,
            nro_nf_gerado INTEGER NOT NULL,
            nro_nf_original INTEGER NOT NULL,
            data_nf TEXT NOT NULL,
            codigo_interno INTEGER NOT NULL,
            causa_devolucao TEXT NOT NULL,
            qtd_devolvida TEXT,
            tipo_unidade_venda TEXT,
            status TEXT NOT NULL DEFAULT 'EXPORTADO',
            exported_at TEXT NOT NULL,
            concluded_at TEXT,
            FOREIGN KEY (lote_id) REFERENCES devolucao_export_lote(id),
            UNIQUE (nro_loja, nro_nf_gerado, data_nf, codigo_interno, causa_devolucao)
        );

        DROP TABLE IF EXISTS integracao_ok_arius;
        CREATE TABLE integracao_ok_arius (
            loja INTEGER NOT NULL,
            dt_gravacao TEXT,
            dt_transmissao TEXT,
            dt_confirmacao TEXT,
            nr_devolucao INTEGER NOT NULL,
            destino_mercadoria TEXT,
            cd_artigo INTEGER NOT NULL,
            artigo_desc TEXT,
            motivo_devolucao TEXT,
            qtd_enviada TEXT,
            qtd_confirmada TEXT,
            importe_preco_custo TEXT,
            status_arius TEXT,
            data_nf TEXT NOT NULL,
            qtd_devolvida TEXT,
            causa_devolucao TEXT NOT NULL,
            status_right TEXT
        );

        CREATE INDEX IF NOT EXISTS idx_integracao_ok_arius_chave
            ON integracao_ok_arius (loja, nr_devolucao, data_nf, cd_artigo, causa_devolucao);

        DROP VIEW IF EXISTS devolucao_export_status_detalhe;
        CREATE VIEW devolucao_export_status_detalhe AS
        SELECT
            e.id,
            l.file_name,
            l.file_path,
            l.exported_at AS lote_exportado_em,
            e.nro_loja,
            e.nro_nf_gerado,
            e.nro_nf_original,
            CASE
                WHEN e.nro_nf_original > 1000 THEN 'MASTER'
                ELSE 'LOJA'
            END AS classificacao_nota,
            e.data_nf,
            e.codigo_interno,
            e.causa_devolucao,
            e.qtd_devolvida,
            e.tipo_unidade_venda,
            e.status,
            e.exported_at,
            e.concluded_at
        FROM devolucao_export_item e
        INNER JOIN devolucao_export_lote l
            ON l.id = e.lote_id;
        """
    )


def normalize_date(value) -> str | None:
    if value in (None, ""):
        return None
    if hasattr(value, "date"):
        return value.date().isoformat()
    return str(value).strip() or None


def normalize_text(value) -> str | None:
    if value in (None, ""):
        return None
    return str(value).strip()


def normalize_int(value) -> int | None:
    if value in (None, ""):
        return None
    return int(float(value))


def normalize_cause(value) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, (int, float)) and float(value).is_integer():
        return str(int(value))
    return str(value).strip()


def derive_original_note(nro_nf_gerado: int) -> int:
    text = str(nro_nf_gerado)
    if len(text) > 4:
        prefix = int(text[:2])
        if 90 <= prefix <= 99:
            return int(text[2:])
    return nro_nf_gerado


def load_rows(input_path: Path, sheet_name: str) -> list[tuple]:
    wb = load_workbook(input_path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows: list[tuple] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        loja = normalize_int(row[0])
        nr_devolucao = normalize_int(row[4])
        cd_artigo = normalize_int(row[6])
        data_nf = normalize_date(row[13])
        causa = normalize_cause(row[15])
        if None in (loja, nr_devolucao, cd_artigo, data_nf):
            continue
        rows.append(
            (
                loja,
                normalize_date(row[1]),
                normalize_date(row[2]),
                normalize_date(row[3]),
                nr_devolucao,
                normalize_text(row[5]),
                cd_artigo,
                normalize_text(row[7]),
                normalize_text(row[8]),
                normalize_text(row[9]),
                normalize_text(row[10]),
                normalize_text(row[11]),
                normalize_text(row[12]),
                data_nf,
                normalize_text(row[14]),
                causa,
                normalize_text(row[16]),
            )
        )
    return rows


def import_ok_rows(conn: sqlite3.Connection, rows: list[tuple]) -> None:
    conn.executemany(
        """
        INSERT INTO integracao_ok_arius (
            loja, dt_gravacao, dt_transmissao, dt_confirmacao, nr_devolucao,
            destino_mercadoria, cd_artigo, artigo_desc, motivo_devolucao,
            qtd_enviada, qtd_confirmada, importe_preco_custo, status_arius,
            data_nf, qtd_devolvida, causa_devolucao, status_right
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        rows,
    )


def mark_concluded(conn: sqlite3.Connection, rows: list[tuple]) -> tuple[int, int]:
    concluded_at = datetime.now().isoformat(timespec="seconds")
    updated = 0
    inserted = 0

    lote = conn.execute(
        """
        INSERT INTO devolucao_export_lote (file_name, file_path, exported_at, total_items)
        VALUES (?, ?, ?, 0)
        ON CONFLICT(file_path) DO UPDATE SET
            file_name = excluded.file_name,
            exported_at = excluded.exported_at
        RETURNING id
        """,
        ("INTEGRACAO_OK_QLIK_ARIUS.xlsx", "/Users/alexandrematheuscrose/Downloads/INTEGRACAO_OK_QLIK_ARIUS.xlsx", concluded_at),
    ).fetchone()
    lote_id = int(lote[0])

    for row in rows:
        loja, _dg, _dt, _dc, nr_devolucao, _dest, cd_artigo, _desc, _mot, _qe, _qc, _ipc, _status, data_nf, qtd_devolvida, causa, _sr = row
        existing = conn.execute(
            """
            SELECT id
            FROM devolucao_export_item
            WHERE nro_loja = ?
              AND nro_nf_gerado = ?
              AND data_nf = ?
              AND codigo_interno = ?
              AND IFNULL(causa_devolucao, '') = ?
            LIMIT 1
            """,
            (loja, nr_devolucao, data_nf, cd_artigo, causa),
        ).fetchone()
        if existing:
            conn.execute(
                """
                UPDATE devolucao_export_item
                SET status = 'CONCLUIDO',
                    concluded_at = ?,
                    lote_id = COALESCE(lote_id, ?)
                WHERE id = ?
                """,
                (concluded_at, lote_id, existing[0]),
            )
            updated += 1
            continue

        conn.execute(
            """
            INSERT OR IGNORE INTO devolucao_export_item (
                lote_id, nro_loja, nro_nf_gerado, nro_nf_original, data_nf,
                codigo_interno, causa_devolucao, qtd_devolvida,
                tipo_unidade_venda, status, exported_at, concluded_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'CONCLUIDO', ?, ?)
            """,
            (
                lote_id,
                loja,
                nr_devolucao,
                derive_original_note(nr_devolucao),
                data_nf,
                cd_artigo,
                causa,
                qtd_devolvida,
                None,
                concluded_at,
                concluded_at,
            ),
        )
        inserted += 1

    conn.execute(
        """
        UPDATE devolucao_export_lote
        SET total_items = (
            SELECT COUNT(*)
            FROM devolucao_export_item
            WHERE lote_id = ?
        )
        WHERE id = ?
        """,
        (lote_id, lote_id),
    )
    return updated, inserted


def main() -> None:
    args = parse_args()
    db_path = Path(args.db).expanduser().resolve()
    input_path = Path(args.input).expanduser().resolve()
    rows = load_rows(input_path, args.sheet)

    with sqlite3.connect(db_path) as conn:
        ensure_schema(conn)
        import_ok_rows(conn, rows)
        updated, inserted = mark_concluded(conn, rows)
        summary = conn.execute(
            """
            SELECT status, COUNT(*)
            FROM devolucao_export_item
            GROUP BY status
            ORDER BY status
            """
        ).fetchall()
        conn.commit()

    print(f"Planilha importada: {input_path}")
    print(f"Banco SQLite: {db_path}")
    print(f"Linhas OK lidas: {len(rows)}")
    print(f"Itens atualizados para concluido: {updated}")
    print(f"Itens inseridos diretamente como concluido: {inserted}")
    for status, total in summary:
        print(f"{status}: {total}")


if __name__ == "__main__":
    main()
