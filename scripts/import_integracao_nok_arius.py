#!/usr/bin/env python3
import argparse
import sqlite3
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Importa a planilha INTEGRACAO_NOK_QLIK_ARIUS e marca devolucoes para reenvio."
    )
    parser.add_argument(
        "--db",
        default="outputs/devolucao_auditoria.sqlite",
        help="Caminho do banco SQLite.",
    )
    parser.add_argument(
        "--input",
        default="/Users/alexandrematheuscrose/Downloads/INTEGRACAO_NOK_QLIK_ARIUS.xlsx",
        help="Caminho do arquivo XLSX.",
    )
    parser.add_argument(
        "--sheet",
        default="COMPARA",
        help="Nome da aba com os registros NOK.",
    )
    return parser.parse_args()


def ensure_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS devolucao_export_lote (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            file_name TEXT NOT NULL,
            file_path TEXT NOT NULL UNIQUE,
            exported_at TEXT NOT NULL,
            total_items INTEGER NOT NULL DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS devolucao_export_item (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            lote_id INTEGER NOT NULL,
            nro_loja INTEGER NOT NULL,
            nro_nf_gerado INTEGER NOT NULL,
            nro_nf_original INTEGER NOT NULL,
            data_nf TEXT NOT NULL,
            codigo_interno INTEGER NOT NULL,
            causa_devolucao TEXT NOT NULL,
            qtd_devolvida TEXT,
            tipo_unidade_venda TEXT,
            status TEXT NOT NULL DEFAULT 'EXPORTADO',
            exported_at TEXT NOT NULL,
            concluded_at TEXT,
            FOREIGN KEY (lote_id) REFERENCES devolucao_export_lote(id),
            UNIQUE (nro_loja, nro_nf_gerado, data_nf, codigo_interno, causa_devolucao)
        );

        DROP TABLE IF EXISTS integracao_nok_arius;
        CREATE TABLE integracao_nok_arius (
            loja_arius TEXT,
            dt_gravacao TEXT,
            dt_transmissao TEXT,
            dt_confirmacao TEXT,
            nr_devolucao_arius TEXT,
            destino_mercadoria TEXT,
            cd_artigo_arius TEXT,
            artigo_desc_arius TEXT,
            motivo_devolucao_arius TEXT,
            qtd_enviada_arius TEXT,
            qtd_confirmada_arius TEXT,
            importe_preco_custo_arius TEXT,
            status_arius TEXT,
            nro_loja INTEGER NOT NULL,
            nro_nf INTEGER NOT NULL,
            data_nf TEXT NOT NULL,
            codigo_interno INTEGER NOT NULL,
            qtd_devolvida TEXT,
            causa_devolucao TEXT NOT NULL,
            tipo_unidade_venda TEXT,
            status_right TEXT
        );

        CREATE INDEX IF NOT EXISTS idx_integracao_nok_arius_chave
            ON integracao_nok_arius (nro_loja, nro_nf, data_nf, codigo_interno, causa_devolucao);

        DROP VIEW IF EXISTS devolucao_export_status_resumo;
        CREATE VIEW devolucao_export_status_resumo AS
        SELECT status, COUNT(*) AS total_itens
        FROM devolucao_export_item
        GROUP BY status;

        DROP VIEW IF EXISTS devolucao_export_status_detalhe;
        CREATE VIEW devolucao_export_status_detalhe AS
        SELECT
            e.id,
            l.file_name,
            l.file_path,
            l.exported_at AS lote_exportado_em,
            e.nro_loja,
            e.nro_nf_gerado,
            e.nro_nf_original,
            CASE
                WHEN e.nro_nf_original > 1000 THEN 'MASTER'
                ELSE 'LOJA'
            END AS classificacao_nota,
            e.data_nf,
            e.codigo_interno,
            e.causa_devolucao,
            e.qtd_devolvida,
            e.tipo_unidade_venda,
            e.status,
            e.exported_at,
            e.concluded_at
        FROM devolucao_export_item e
        INNER JOIN devolucao_export_lote l
            ON l.id = e.lote_id;
        """
    )


def normalize_date(value) -> str | None:
    if value in (None, ""):
        return None
    if hasattr(value, "date"):
        return value.date().isoformat()
    text = str(value).strip()
    if not text:
        return None
    return text


def normalize_text(value) -> str | None:
    if value in (None, ""):
        return None
    return str(value).strip()


def normalize_int(value) -> int | None:
    if value in (None, ""):
        return None
    return int(float(value))


def normalize_cause(value) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, (int, float)) and float(value).is_integer():
        return str(int(value))
    return str(value).strip()


def derive_original_note(nro_nf_gerado: int) -> int:
    text = str(nro_nf_gerado)
    if len(text) > 4:
        prefix = int(text[:2])
        if 90 <= prefix <= 99:
            return int(text[2:])
    return nro_nf_gerado


def classify_note(nro_nf_original: int) -> str:
    return "MASTER" if int(nro_nf_original) > 1000 else "LOJA"


def load_rows(input_path: Path, sheet_name: str) -> list[tuple]:
    wb = load_workbook(input_path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows: list[tuple] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        nro_loja = normalize_int(row[13])
        nro_nf = normalize_int(row[14])
        data_nf = normalize_date(row[15])
        codigo_interno = normalize_int(row[16])
        causa = normalize_cause(row[18])
        if None in (nro_loja, nro_nf, data_nf, codigo_interno):
            continue
        rows.append(
            (
                normalize_text(row[0]),
                normalize_date(row[1]),
                normalize_date(row[2]),
                normalize_date(row[3]),
                normalize_text(row[4]),
                normalize_text(row[5]),
                normalize_text(row[6]),
                normalize_text(row[7]),
                normalize_text(row[8]),
                normalize_text(row[9]),
                normalize_text(row[10]),
                normalize_text(row[11]),
                normalize_text(row[12]),
                nro_loja,
                nro_nf,
                data_nf,
                codigo_interno,
                normalize_text(row[17]),
                causa,
                normalize_text(row[19]),
                normalize_text(row[20]),
            )
        )
    return rows


def get_or_create_reenvio_lote(conn: sqlite3.Connection, input_path: Path) -> int:
    exported_at = datetime.now().isoformat(timespec="seconds")
    lote = conn.execute(
        """
        INSERT INTO devolucao_export_lote (file_name, file_path, exported_at, total_items)
        VALUES (?, ?, ?, 0)
        ON CONFLICT(file_path) DO UPDATE SET
            file_name = excluded.file_name,
            exported_at = excluded.exported_at
        RETURNING id
        """,
        (input_path.name, str(input_path), exported_at),
    ).fetchone()
    return int(lote[0])


def import_nok_rows(conn: sqlite3.Connection, rows: list[tuple]) -> None:
    conn.executemany(
        """
        INSERT INTO integracao_nok_arius (
            loja_arius, dt_gravacao, dt_transmissao, dt_confirmacao, nr_devolucao_arius,
            destino_mercadoria, cd_artigo_arius, artigo_desc_arius, motivo_devolucao_arius,
            qtd_enviada_arius, qtd_confirmada_arius, importe_preco_custo_arius, status_arius,
            nro_loja, nro_nf, data_nf, codigo_interno, qtd_devolvida, causa_devolucao,
            tipo_unidade_venda, status_right
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        rows,
    )


def mark_for_reenvio(conn: sqlite3.Connection, lote_id: int, rows: list[tuple]) -> tuple[int, int]:
    inserted = 0
    updated = 0
    exported_at = datetime.now().isoformat(timespec="seconds")

    for row in rows:
        nro_loja = row[13]
        nro_nf = row[14]
        data_nf = row[15]
        codigo_interno = row[16]
        qtd_devolvida = row[17]
        causa = row[18] or ""
        tipo_unidade = row[19]
        nro_nf_original = derive_original_note(nro_nf)
        target_status = "MASTER" if classify_note(nro_nf_original) == "MASTER" else "REENVIAR"

        existing = conn.execute(
            """
            SELECT id, status
            FROM devolucao_export_item
            WHERE nro_loja = ?
              AND nro_nf_gerado = ?
              AND data_nf = ?
              AND codigo_interno = ?
              AND IFNULL(causa_devolucao, '') = ?
            LIMIT 1
            """,
            (nro_loja, nro_nf, data_nf, codigo_interno, causa),
        ).fetchone()

        if existing:
            conn.execute(
                """
                UPDATE devolucao_export_item
                SET status = ?,
                    concluded_at = NULL
                WHERE id = ?
                """,
                (target_status, existing[0]),
            )
            updated += 1
            continue

        conn.execute(
            """
            INSERT INTO devolucao_export_item (
                lote_id, nro_loja, nro_nf_gerado, nro_nf_original, data_nf,
                codigo_interno, causa_devolucao, qtd_devolvida,
                tipo_unidade_venda, status, exported_at, concluded_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, NULL)
            """,
            (
                lote_id,
                nro_loja,
                nro_nf,
                nro_nf_original,
                data_nf,
                codigo_interno,
                causa,
                qtd_devolvida,
                tipo_unidade,
                target_status,
                exported_at,
            ),
        )
        inserted += 1

    conn.execute(
        """
        UPDATE devolucao_export_lote
        SET total_items = (
            SELECT COUNT(*)
            FROM devolucao_export_item
            WHERE lote_id = ?
        )
        WHERE id = ?
        """,
        (lote_id, lote_id),
    )
    return inserted, updated


def main() -> None:
    args = parse_args()
    db_path = Path(args.db).expanduser().resolve()
    input_path = Path(args.input).expanduser().resolve()

    rows = load_rows(input_path, args.sheet)
    with sqlite3.connect(db_path) as conn:
        ensure_schema(conn)
        import_nok_rows(conn, rows)
        lote_id = get_or_create_reenvio_lote(conn, input_path)
        inserted, updated = mark_for_reenvio(conn, lote_id, rows)
        summary = conn.execute(
            """
            SELECT status, COUNT(*)
            FROM devolucao_export_item
            GROUP BY status
            ORDER BY status
            """
        ).fetchall()
        conn.commit()

    print(f"Planilha importada: {input_path}")
    print(f"Banco SQLite: {db_path}")
    print(f"Linhas NOK lidas: {len(rows)}")
    print(f"Itens inseridos para reenvio: {inserted}")
    print(f"Itens atualizados para reenvio: {updated}")
    for status, total in summary:
        print(f"{status}: {total}")


if __name__ == "__main__":
    main()
