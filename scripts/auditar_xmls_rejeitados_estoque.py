#!/usr/bin/env python3
"""Audita XMLs rejeitados contra notas e movimentos de estoque do Pleno.

Consulta somente leitura. Um item e considerado movimentado apenas quando ha
registro em EST05 vinculado diretamente ao item fiscal (FIS02) da nota.
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import sys
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / ".python_packages"))


def load_env(path: Path) -> dict[str, str]:
    values: dict[str, str] = {}
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if line and not line.startswith("#") and "=" in line:
            key, value = line.split("=", 1)
            values[key.strip()] = value.strip().strip('"').strip("'")
    values.update(os.environ)
    return values


def decimal(value: str | None) -> Decimal:
    try:
        return Decimal((value or "0").strip())
    except (InvalidOperation, AttributeError):
        return Decimal("0")


def clean_cnpj(value: str | None) -> str:
    return re.sub(r"\D", "", value or "")


def chunked(values: list[object], size: int = 700):
    for index in range(0, len(values), size):
        yield values[index:index + size]


def parse_xml(path: Path) -> dict[str, object] | None:
    try:
        root = ET.parse(path).getroot()
    except ET.ParseError:
        return None

    def first(name: str) -> str:
        node = root.find(f".//{{*}}{name}")
        return (node.text or "").strip() if node is not None else ""

    inf = root.find(".//{*}infNFe")
    key = first("chNFe")
    if not key and inf is not None:
        key = (inf.attrib.get("Id") or "").replace("NFe", "")
    issued = first("dhEmi") or first("dEmi")
    issued_day = issued[:10]
    items = []
    for detail in root.findall(".//{*}det"):
        product = detail.find("{*}prod")
        if product is None:
            continue
        get = lambda name: ((product.findtext(f"{{*}}{name}") or "").strip())
        qty = decimal(get("qCom"))
        total = decimal(get("vProd"))
        items.append({
            "item_xml": detail.attrib.get("nItem", ""),
            "codigo_produto": get("cProd"),
            "descricao_produto": get("xProd"),
            "quantidade_xml": qty,
            "valor_xml": total,
        })
    return {
        "arquivo_xml": str(path),
        "chave": key,
        "nota": first("nNF"),
        "serie": first("serie"),
        "emissao": issued_day,
        "emitente_cnpj": clean_cnpj(first("emit/{*}CNPJ") or first("CNPJ")),
        "destinatario_cnpj": clean_cnpj(first("dest/{*}CNPJ")),
        "destinatario_nome": first("dest/{*}xNome"),
        "itens": items,
    }


def query_filiais(conn) -> dict[str, dict[str, object]]:
    sql = """
        SELECT DISTINCT
            REPLACE(REPLACE(REPLACE(e.pes03_cnpj, '.', ''), '/', ''), '-', '') AS cnpj,
            f.cfg06_numero AS loja,
            f.cfg06_nome AS nome_loja
        FROM cfg06_filial f
        JOIN pes03_estabelecimento e ON e.pes03_id = f.pes03_estabelecimento_id
        WHERE e.pes03_cnpj IS NOT NULL
          AND f.cfg06_numero IS NOT NULL
    """
    with conn.cursor() as cur:
        cur.execute(sql)
        rows = cur.fetchall()
    result = {}
    for row in rows:
        cnpj = clean_cnpj(str(row["cnpj"] or ""))
        loja = row["loja"]
        if cnpj and loja is not None:
            # Prefer active-looking smaller store number when CNPJ has duplicates.
            previous = result.get(cnpj)
            if previous is None or int(loja) < int(previous["loja"]):
                result[cnpj] = {"loja": int(loja), "nome_loja": row["nome_loja"] or ""}
    return result


def query_pleno_items(conn, keys: list[str]) -> tuple[dict[str, list[dict]], dict[int, dict]]:
    notes: dict[str, list[dict]] = defaultdict(list)
    items_by_id: dict[int, dict] = {}
    for values in chunked(keys):
        placeholders = ",".join(["%s"] * len(values))
        sql = f"""
            SELECT
                nf.fis01_id AS nota_id,
                nf.fis01_nfe_chave AS chave,
                nf.fis01_nronf AS nota_pleno,
                nf.fis01_serienf AS serie_pleno,
                nf.fis01_data_emissao AS emissao_pleno,
                nf.fis01_data_entrada_saida AS entrada_pleno,
                CASE WHEN nf.adm05_usuario_checkin_id IS NULL THEN 'SEM_CHECKIN' ELSE 'COM_CHECKIN' END AS checkin,
                nf.fis01_flgcancelada AS cancelada,
                fd.cfg06_numero AS loja_pleno,
                fd.cfg06_nome AS nome_loja_pleno,
                ni.fis02_id AS item_id,
                ni.fis02_qtd AS quantidade_pleno,
                ni.fis02_vlrunit AS valor_unitario_pleno,
                ni.fis02_vlrunit_liquido AS valor_unitario_liquido,
                m.mcd01_codint AS codigo_produto_pleno,
                m.mcd01_descricao AS descricao_produto_pleno
            FROM fis01_notafiscal nf
            LEFT JOIN fis02_notafiscal_item ni ON ni.fis01_notafiscal_id = nf.fis01_id
            LEFT JOIN mcd03_mercadoria_filial mf ON mf.mcd03_id = ni.mcd03_mercadoria_filial_id
            LEFT JOIN mcd01_mercadoria m ON m.mcd01_id = mf.mcd01_mercadoria_id
            LEFT JOIN pes04_pessoa pd ON pd.pes04_id = nf.pes04_pessoa_dest
            LEFT JOIN pes03_estabelecimento ed ON ed.pes03_id = pd.pes03_estabelecimento_id
            LEFT JOIN cfg06_filial fd ON fd.pes03_estabelecimento_id = ed.pes03_id
            WHERE nf.fis01_nfe_chave IN ({placeholders})
            ORDER BY nf.fis01_id, ni.fis02_id
        """
        with conn.cursor() as cur:
            cur.execute(sql, values)
            rows = cur.fetchall()
        for row in rows:
            key = str(row["chave"] or "")
            notes[key].append(row)
            if row["item_id"] is not None:
                items_by_id[int(row["item_id"])] = row

    movements: dict[int, dict] = {}
    ids = list(items_by_id)
    for values in chunked(ids):
        placeholders = ",".join(["%s"] * len(values))
        sql = f"""
            SELECT
                fis02_notafiscal_item_id AS item_id,
                COUNT(*) AS movimentos_estoque,
                MIN(est05_data) AS primeira_data_movimento,
                MAX(est05_data) AS ultima_data_movimento
            FROM est05_estoque_movimento
            WHERE fis02_notafiscal_item_id IN ({placeholders})
            GROUP BY fis02_notafiscal_item_id
        """
        with conn.cursor() as cur:
            cur.execute(sql, values)
            movements.update({int(row["item_id"]): row for row in cur.fetchall()})
    return notes, movements


def write_csv(path: Path, rows: list[dict], fields: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as output:
        writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def autosize(sheet) -> None:
    from openpyxl.utils import get_column_letter
    for column in range(1, sheet.max_column + 1):
        width = min(max(len(str(sheet.cell(row, column).value or "")) for row in range(1, min(sheet.max_row, 250) + 1)) + 2, 54)
        sheet.column_dimensions[get_column_letter(column)].width = max(width, 12)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--start-date", default="2026-08-01")
    parser.add_argument("--end-date", default=date.today().isoformat())
    args = parser.parse_args()
    start = date.fromisoformat(args.start_date)
    end = date.fromisoformat(args.end_date)

    import pymysql
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    env = load_env(ROOT / ".env")
    conn = pymysql.connect(
        host=env["MYSQL_HOST"], port=int(env.get("MYSQL_PORT", "3306")),
        user=env["MYSQL_USER"], password=env["MYSQL_PASSWORD"], database=env["MYSQL_DATABASE"],
        charset="latin1", cursorclass=pymysql.cursors.DictCursor,
    )
    try:
        filiais = query_filiais(conn)
        xmls = []
        invalid = 0
        for path in sorted((ROOT / "rejeitados").rglob("*.xml")):
            record = parse_xml(path)
            if record is None:
                invalid += 1
                continue
            emitted = record["emissao"]
            try:
                emitted_date = date.fromisoformat(str(emitted))
            except ValueError:
                continue
            store = filiais.get(str(record["destinatario_cnpj"]))
            if not store or not (0 < store["loja"] < 3000) or store["loja"] == 704:
                continue
            record.update(store)
            if start <= emitted_date <= end:
                xmls.append(record)
        # The same XML may occur more than once in the rejected folder.
        unique = {str(row["chave"]): row for row in xmls if row["chave"]}
        xmls = list(unique.values())
        keys = list(unique)
        pleno_by_key, movements = query_pleno_items(conn, keys)
    finally:
        conn.close()

    note_rows: list[dict] = []
    item_rows: list[dict] = []
    rebuild: dict[tuple[object, str], dict] = {}
    for xml in sorted(xmls, key=lambda value: (value["loja"], value["emissao"], value["nota"])):
        pleno = pleno_by_key.get(str(xml["chave"]), [])
        pleno_items = [row for row in pleno if row["item_id"] is not None]
        by_product: dict[str, list[dict]] = defaultdict(list)
        for row in pleno_items:
            by_product[str(row["codigo_produto_pleno"] or "")].append(row)
        moved_count = 0
        no_move_count = 0
        missing_count = 0
        for item in xml["itens"]:
            candidates = by_product.get(str(item["codigo_produto"]), [])
            selected = candidates[0] if candidates else None
            movement = movements.get(int(selected["item_id"])) if selected else None
            movement_count = int(movement["movimentos_estoque"]) if movement else 0
            if selected is None:
                item_status = "ITEM_NAO_ENCONTRADO_NO_PLENO" if pleno else "NOTA_NAO_ENCONTRADA_NO_PLENO"
                missing_count += 1
            elif movement_count:
                item_status = "MOVIMENTOU_ESTOQUE"
                moved_count += 1
            else:
                item_status = "SEM_MOVIMENTACAO_ESTOQUE"
                no_move_count += 1
            unit = decimal(str((selected or {}).get("valor_unitario_liquido") or (selected or {}).get("valor_unitario_pleno") or 0))
            row = {
                "loja": xml["loja"], "nome_loja": xml["nome_loja"], "chave": xml["chave"], "nota_xml": xml["nota"],
                "serie_xml": xml["serie"], "emissao_xml": xml["emissao"], "item_xml": item["item_xml"],
                "codigo_produto_xml": item["codigo_produto"], "descricao_produto_xml": item["descricao_produto"],
                "quantidade_xml": float(item["quantidade_xml"]), "valor_xml": float(item["valor_xml"]),
                "status_item": item_status, "nota_id_pleno": (selected or {}).get("nota_id"),
                "nota_pleno": (selected or {}).get("nota_pleno"), "checkin_pleno": (selected or {}).get("checkin"),
                "codigo_produto_pleno": (selected or {}).get("codigo_produto_pleno"),
                "descricao_produto_pleno": (selected or {}).get("descricao_produto_pleno"),
                "quantidade_pleno": float(decimal(str((selected or {}).get("quantidade_pleno") or 0))),
                "valor_pleno": float(decimal(str((selected or {}).get("quantidade_pleno") or 0)) * unit),
                "movimentos_estoque": movement_count,
                "primeira_data_movimento": (movement or {}).get("primeira_data_movimento"),
                "ultima_data_movimento": (movement or {}).get("ultima_data_movimento"),
            }
            item_rows.append(row)
            if item_status != "MOVIMENTOU_ESTOQUE":
                key = (xml["loja"], str(item["codigo_produto"]))
                target = rebuild.setdefault(key, {"loja": xml["loja"], "nome_loja": xml["nome_loja"], "codigo_produto": item["codigo_produto"], "descricao_produto": item["descricao_produto"], "quantidade_total": Decimal("0"), "valor_total": Decimal("0"), "notas": set(), "motivos": set()})
                target["quantidade_total"] += item["quantidade_xml"]
                target["valor_total"] += item["valor_xml"]
                target["notas"].add(str(xml["nota"]))
                target["motivos"].add(item_status)
        if not pleno:
            status_note = "NOTA_NAO_ENTROU_NO_PLENO"
        elif not pleno_items:
            status_note = "NOTA_ENTROU_SEM_ITENS_NO_PLENO"
        elif moved_count == len(xml["itens"]):
            status_note = "NOTA_E_ITENS_MOVIMENTARAM_ESTOQUE"
        elif moved_count:
            status_note = "MOVIMENTACAO_PARCIAL_DE_ITENS"
        else:
            status_note = "NOTA_SEM_MOVIMENTACAO_DE_ESTOQUE"
        example = pleno[0] if pleno else {}
        note_rows.append({
            "status_nota": status_note, "loja": xml["loja"], "nome_loja": xml["nome_loja"], "chave": xml["chave"],
            "nota_xml": xml["nota"], "serie_xml": xml["serie"], "emissao_xml": xml["emissao"],
            "itens_xml": len(xml["itens"]), "quantidade_xml": float(sum((item["quantidade_xml"] for item in xml["itens"]), Decimal("0"))),
            "valor_xml": float(sum((item["valor_xml"] for item in xml["itens"]), Decimal("0"))),
            "nota_id_pleno": example.get("nota_id"), "nota_pleno": example.get("nota_pleno"), "checkin_pleno": example.get("checkin"),
            "itens_movimentaram": moved_count, "itens_sem_movimento": no_move_count, "itens_nao_encontrados": missing_count,
            "movimentos_estoque": sum(int(movements.get(int(row["item_id"]), {}).get("movimentos_estoque", 0)) for row in pleno_items),
        })

    grouped_notes: dict[str, list[dict]] = defaultdict(list)
    for note in note_rows:
        grouped_notes[note["status_nota"]].append(note)
    summary = []
    for status, grouped in sorted(grouped_notes.items()):
        summary.append({"grupo": status, "notas": len(grouped), "itens_xml": sum(row["itens_xml"] for row in grouped), "quantidade_xml": round(sum(row["quantidade_xml"] for row in grouped), 3), "valor_xml": round(sum(row["valor_xml"] for row in grouped), 2)})
    summary.insert(0, {"grupo": "TOTAL_XMLS_REJEITADOS_AUDITADOS", "notas": len(note_rows), "itens_xml": len(item_rows), "quantidade_xml": round(sum(row["quantidade_xml"] for row in note_rows), 3), "valor_xml": round(sum(row["valor_xml"] for row in note_rows), 2)})
    rebuild_rows = [{**row, "quantidade_total": float(row["quantidade_total"]), "valor_total": float(row["valor_total"]), "notas": ", ".join(sorted(row["notas"], key=int)), "motivos": ", ".join(sorted(row["motivos"]))} for row in rebuild.values()]
    rebuild_rows.sort(key=lambda row: (row["loja"], row["codigo_produto"]))

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = ROOT / "outputs" / "rejeitados_lojas_proprias"
    out.mkdir(parents=True, exist_ok=True)
    base = out / f"auditoria_xmls_rejeitados_pleno_estoque_{args.start_date}_a_{args.end_date}_{stamp}"
    note_fields = list(note_rows[0]) if note_rows else []
    item_fields = list(item_rows[0]) if item_rows else []
    rebuild_fields = ["loja", "nome_loja", "codigo_produto", "descricao_produto", "quantidade_total", "valor_total", "notas", "motivos"]
    write_csv(base.with_name(base.name + "_notas.csv"), note_rows, note_fields)
    write_csv(base.with_name(base.name + "_itens.csv"), item_rows, item_fields)
    write_csv(base.with_name(base.name + "_itens_sem_movimento_agregados.csv"), rebuild_rows, rebuild_fields)
    write_csv(base.with_name(base.name + "_resumo.csv"), summary, ["grupo", "notas", "itens_xml", "quantidade_xml", "valor_xml"])

    book = Workbook()
    book.remove(book.active)
    for title, rows, fields in [("Resumo", summary, ["grupo", "notas", "itens_xml", "quantidade_xml", "valor_xml"]), ("Notas", note_rows, note_fields), ("Itens", item_rows, item_fields), ("Sem movimento", rebuild_rows, rebuild_fields)]:
        sheet = book.create_sheet(title)
        sheet.append(fields)
        for row in rows:
            sheet.append([row.get(field, "") for field in fields])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        autosize(sheet)
    xlsx = base.with_suffix(".xlsx")
    book.save(xlsx)
    print(f"xmls_lidos={len(xmls)} invalidos={invalid} notas={len(note_rows)} itens={len(item_rows)}")
    print(f"planilha={xlsx}")
    print(f"resumo={base.with_name(base.name + '_resumo.csv')}")


if __name__ == "__main__":
    main()
